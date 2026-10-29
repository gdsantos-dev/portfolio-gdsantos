"""Generate the editable, one-page Black & Crimson curriculum from source."""
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

OUT = Path(__file__).resolve().parents[1] / 'public/data/curriculos/guilherme-dos-santos-curriculo-pt.docx'
FONT = 'Arial'
INK = RGBColor(29, 29, 32)
CRIMSON = RGBColor(204, 27, 56)
MUTED = RGBColor(88, 88, 94)
WHITE = RGBColor(255, 255, 255)

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        tcPr.append(shd)
    shd.set(qn('w:fill'), fill)

def set_cell_margin(cell, top=0, left=0, bottom=0, right=0):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in('w:tcMar')
    if tcMar is None:
        tcMar = OxmlElement('w:tcMar')
        tcPr.append(tcMar)
    for margin, value in [('top',top),('start',left),('bottom',bottom),('end',right)]:
        node = tcMar.find(qn('w:' + margin))
        if node is None:
            node = OxmlElement('w:' + margin)
            tcMar.append(node)
        node.set(qn('w:w'), str(value))
        node.set(qn('w:type'), 'dxa')

def set_border_bottom(paragraph, color='E3CFD4', size='6'):
    pPr = paragraph._p.get_or_add_pPr()
    borders = pPr.find(qn('w:pBdr'))
    if borders is None:
        borders = OxmlElement('w:pBdr')
        pPr.append(borders)
    bot = OxmlElement('w:bottom')
    for name, val in [('val','single'),('sz',size),('color',color),('space','2')]:
        bot.set(qn('w:' + name), val)
    borders.append(bot)

def no_borders(table):
    tblPr = table._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        element = OxmlElement('w:' + side)
        element.set(qn('w:val'), 'nil')
        borders.append(element)
    tblPr.append(borders)

def add_run(p, text, size=9.8, bold=False, color=INK):
    r = p.add_run(text)
    r.bold = bold
    r.font.name = FONT
    r.font.size = Pt(size)
    r.font.color.rgb = color
    fonts = r._element.get_or_add_rPr().rFonts
    if fonts is not None:
        fonts.set(qn('w:ascii'), FONT)
        fonts.set(qn('w:hAnsi'), FONT)
    return r

def add_link(p, label, url, color=CRIMSON, size=9.4):
    relation = p.part.relate_to(url, RT.HYPERLINK, is_external=True)
    h = OxmlElement('w:hyperlink')
    h.set(qn('r:id'), relation)
    r = OxmlElement('w:r')
    rPr = OxmlElement('w:rPr')
    col = OxmlElement('w:color')
    col.set(qn('w:val'), '%02X%02X%02X' % (color[0], color[1], color[2]))
    rPr.append(col)
    sz = OxmlElement('w:sz')
    sz.set(qn('w:val'), str(int(round(size*2))))
    rPr.append(sz)
    rFonts = OxmlElement('w:rFonts')
    rFonts.set(qn('w:ascii'), FONT)
    rFonts.set(qn('w:hAnsi'), FONT)
    rPr.append(rFonts)
    r.append(rPr)
    t = OxmlElement('w:t')
    t.text = label
    r.append(t)
    h.append(r)
    p._p.append(h)

def space_p(p, before=0, after=0, line=1.2):
    fmt = p.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line

def section_heading(doc, text, before=13):
    p = doc.add_paragraph(style='SectionLabel')
    space_p(p, before, 6, 1.05)
    p.paragraph_format.keep_with_next = True
    add_run(p, text, 10.2, True, CRIMSON)
    set_border_bottom(p)
    return p

def bullet(doc, text):
    p = doc.add_paragraph(style='CustomBullet')
    fmt = p.paragraph_format
    fmt.left_indent = Inches(.15)
    fmt.first_line_indent = Inches(-.15)
    space_p(p, 2.8, 1.5, 1.22)
    add_run(p, '•', 9.6, False, CRIMSON)
    add_run(p, '  ' + text, 9.5, False, INK)
    return p

d = Document()
sec = d.sections[0]
sec.page_width = Inches(8.27)
sec.page_height = Inches(11.69)
sec.left_margin = Inches(.68)
sec.right_margin = Inches(.68)
sec.top_margin = Inches(.55)
sec.bottom_margin = Inches(.57)
sec.header_distance = Inches(.25)
sec.footer_distance = Inches(.35)
d.core_properties.title = 'Currículo | Guilherme dos Santos'
d.core_properties.subject = 'Currículo profissional'
d.core_properties.author = 'Guilherme dos Santos'

styles = d.styles
normal = styles['Normal']
normal.font.name = FONT
normal.font.size = Pt(9.6)
normal.font.color.rgb = INK
normal.paragraph_format.space_before = Pt(0)
normal.paragraph_format.space_after = Pt(0)
for name in ['SectionLabel','CustomBullet']:
    st = styles[name] if name in styles else styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    st.font.name = FONT
    st.font.size = Pt(9.6)

band = d.add_table(rows=1, cols=1)
band.autofit = False
band.columns[0].width = Inches(6.9)
cell = band.cell(0,0)
cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
shade(cell, '0D0D0F')
set_cell_margin(cell, top=260, left=265, bottom=250, right=265)
no_borders(band)
p = cell.paragraphs[0]
space_p(p,0,4,1.1)
add_run(p,'GUILHERME DOS SANTOS',20,True,WHITE)
p = cell.add_paragraph()
space_p(p,0,0,1.1)
add_run(p,'ESTUDANTE DE CIÊNCIA DA COMPUTAÇÃO',9.7,True,RGBColor(255,81,105))

p = d.add_paragraph()
space_p(p,9,3,1.1)
add_run(p,'Recife e Região Metropolitana, PE',9.2,False,MUTED)
add_run(p,'   ·   ',9.2,False,MUTED)
add_link(p,'(81) 99597-5655','tel:+5581995975655',size=9.2)
p = d.add_paragraph()
space_p(p,0,0,1.1)
add_link(p,'guilhermesantos_26@hotmail.com','mailto:guilhermesantos_26@hotmail.com',size=9.2)
add_run(p,'   ·   ',9.2,False,MUTED)
add_link(p,'GitHub','https://github.com/gdsantos-dev',size=9.2)
add_run(p,'   ·   ',9.2,False,MUTED)
add_link(p,'LinkedIn','https://www.linkedin.com/in/guilherme-santos-b4169740b/',size=9.2)

section_heading(d,'OBJETIVO PROFISSIONAL',before=14)
p=d.add_paragraph()
space_p(p,1,0,1.29)
add_run(p,'Busco estágio em desenvolvimento de software, com interesse em aplicações web (front-end, back-end ou full stack), engenharia de software e inteligência artificial aplicada a produtos digitais.',9.6)

section_heading(d,'FORMAÇÃO ACADÊMICA',before=15)
p=d.add_paragraph()
space_p(p,1,3,1.17)
add_run(p,'Bacharelado em Ciência da Computação',9.8,True)
add_run(p,'   |   ',9.8,False,MUTED)
add_run(p,'UniFG',9.8,False,MUTED)
p=d.add_paragraph()
space_p(p,0,0,1.16)
add_run(p,'6º período em 2026   ·   Conclusão prevista: dezembro de 2027',9.3,False,MUTED)

section_heading(d,'ATIVIDADES PRÁTICAS E ACADÊMICAS',before=16)
p=d.add_paragraph()
space_p(p,1,3,1.2)
add_run(p,'Desenvolvimento web e prototipação',9.6,True)
add_run(p,'   |   Estudos e atividades pessoais · 2026',9.2,False,MUTED)
bullet(d,'Criação e refinamento de interfaces responsivas, navegação e fluxos de interação em protótipos web.')
bullet(d,'Prática com HTML, CSS e JavaScript, incluindo leitura, ajustes e depuração de código.')
p=d.add_paragraph()
space_p(p,7,3,1.2)
add_run(p,'Organização e validação de soluções',9.6,True)
add_run(p,'   |   Estudos e atividades acadêmicas',9.2,False,MUTED)
bullet(d,'Documentação de requisitos, análise de usabilidade e execução de testes funcionais.')
bullet(d,'Pesquisa e comparação de alternativas técnicas, com ferramentas de IA como apoio ao aprendizado e à implementação.')

section_heading(d,'HABILIDADES TÉCNICAS',before=15)
for label, values in [
    ('Web','HTML, CSS, JavaScript; contato com TypeScript'),
    ('Back-end e dados','Python, FastAPI, SQLite; contato com Django'),
    ('Ferramentas e práticas','Git, GitHub, APIs REST, prototipação, testes funcionais e documentação'),
]:
    p=d.add_paragraph()
    space_p(p,2.2,4.2,1.17)
    p.paragraph_format.keep_together = True
    add_run(p,label+': ',9.6,True,CRIMSON)
    add_run(p,values,9.5,False,INK)

section_heading(d,'IDIOMAS',before=12)
p=d.add_paragraph()
space_p(p,2,0,1.17)
add_run(p,'Inglês',9.7,True,INK)
add_run(p,'   ·   Intermediário',9.5,False,MUTED)

footer=sec.footer.paragraphs[0]
footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
space_p(footer,0,0,1)
add_run(footer,'Guilherme dos Santos   ·   Currículo profissional',8,False,MUTED)

OUT.parent.mkdir(parents=True,exist_ok=True)
d.save(OUT)
print(OUT)
